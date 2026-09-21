# ShareSansar corporate-event calendar → formatted Excel → email
# GitHub Actions edition of Sharesansar_Daily_event_download_v2.ipynb
#
# Behaviour requested:
#   * Runs 08:30 NPT, Monday–Friday (scheduled via the workflow cron).
#   * Scrapes the NUM_DAYS=3 days starting 3 days before the run date (today-3 .. today-1).
#   * Emails the formatted .xlsx as an attachment from Gmail to the office acct.
#
# Uses plain requests (no Selenium), so it runs fast and reliably on a runner.
# Scraping / formatting logic is unchanged from the notebook.

import os
import re
import ssl
import smtplib
import tempfile
from email.message import EmailMessage
from datetime import datetime, timedelta, timezone

import requests
import pandas as pd
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ----------------- CONFIG -----------------
# NPT = UTC+5:45. Compute "today" in Nepal time, then look 3 days before
NPT = timezone(timedelta(hours=5, minutes=45))
TODAY_NPT = datetime.now(NPT).replace(hour=0, minute=0, second=0, microsecond=0, tzinfo=None)
START_DATE = TODAY_NPT - timedelta(days=3)
NUM_DAYS = 3
SAVE_DIR = tempfile.gettempdir()      # runner temp; file only lives in the email
ALSO_SAVE_CSV = False

# ----------------- MAIL -----------------
GMAIL_USER = os.environ.get("GMAIL_USER", "")
GMAIL_APP_PW = os.environ.get("GMAIL_APP_PW", "")
MAIL_TO = os.environ.get("MAIL_TO", "")
SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 465
# -------------------------------------------

# ---------- Scraper (resilient: one bad page won't kill the run) ----------
session = requests.Session()
session.headers.update({"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})


def fetch_soup(url, timeout=20):
    try:
        r = session.get(url, timeout=timeout)
        r.raise_for_status()
        return BeautifulSoup(r.text, "lxml")
    except requests.RequestException as e:
        print(f"  ! skipped {url}: {e}")
        return None


def fetch_detail(url):
    soup = fetch_soup(url)
    if soup is None:
        return ""
    div = soup.find("div", {"class": "detail b-shadow margin-bottom-20"})
    return div.get_text("\n", strip=True) if div else ""


def scrape():
    rows = []
    for i in range(NUM_DAYS):
        day = START_DATE + timedelta(days=i)
        url = "https://www.sharesansar.com/events/" + day.strftime("%Y-%m-%d")
        print(f"Fetching {url}")
        soup = fetch_soup(url)
        if soup is None:
            continue
        for item in soup.find_all("div", {"class": "featured-news-list margin-bottom-15"}):
            d = item.find("span", {"class": "text-org"})
            el = item.find("h4", {"class": "featured-news-title"})
            a = item.find("a")
            link = a["href"] if a and a.has_attr("href") else ""
            rows.append({
                "Date": d.text.strip() if d else "N/A",
                "Event": el.text.strip() if el else "N/A",
                "Link": link,
                "Details": fetch_detail(link) if link else "",
            })
    return rows


# ---------- Clean, classify, dedupe ----------
# First match wins, so order matters (e.g. Book Closure before AGM/Dividend).
TYPE_RULES = [
    ("Book Closure", r"bl?ook\s*c?lo[sc]"),   # also typos "Book losure", "Book Clocure", "Blook Close"
    ("AGM", r"(?<![a-z])agm(?![a-z])|annual general meeting|[बव]ा\S*क साधारण"),
    ("SGM", r"(?<![a-z])(sgm|egm|eogm)(?![a-z])|s\.g\.m|spe\w*ial\s+general|extra\s*-?ordinary general|"
            r"विशेष साधारण सभा"),
    ("Auction", r"auction"),
    ("IPO", r"\bipo\b|\bfpo\b"),
    ("Listing", r"\blisting\b"),
    ("Right Share", r"rights? share|\bright from\b|issuance of right|हकप्रद"),
    ("Dividend", r"div[ie]dend|bonus|लाभांश|लाभांस"),
    ("Preference Share", r"preference share|अग्राधिकार"),
    ("Lock-in", r"lock[- ]?(in|end)|लक-?ईन"),
    ("Promoter Conversion", r"conversion"),
    ("Debenture/Bond", r"debenture|\bbond\b|redemption"),
    # Generic issue windows (FPI/local/migrant/mutual fund units): after the
    # specific issue types above so their opening/closing days stay with them.
    ("IPO", r"opening day|closing day"),
    ("Suspension/Delisting", r"suspend|delist"),
    ("Merger/Acquisition", r"merger|मर्जर|\bmou\b|acqui|integration|joint (operation|transaction|banking)|swap ratio"),
    ("Share Sale", r"sale of shares?|share sal|share sell|शेयर बिक्री"),
    ("Share Freeze/Pledge", r"rokka|freez|pledge|रोक्का|धितोमा"),
    ("Price Adjustment", r"price adjust"),
    ("Commercial Operation", r"commercial operat|\bcod\b|commissioning"),
    ("Plant/Disaster", r"shut\s*down|resum|stoppage|outage|plant unit|plant (operation|partially)|in operation|"
                       r"reoperation|force majeure|flood|damage|बाढी|विपत्ति|बिशेष परिस्थिति|विद्युत उत्पादन|"
                       r"आयोजना बन्द"),
    ("Financial Report", r"quarter|quaterly|\bqtr\b|financial (report|statement)|annual report|audit report|"
                         r"governance report"),
    ("Credit Rating", r"\brating"),
    ("RTS/Registrar", r"\brts\b|registrar|share register|रजिष्टर|\bisin\b"),
    ("Auditor", r"\baudi?t?ors?\b(?!\s*committee)|statutory audit|लेखा ?पर[ीि]क्षक"),   # also "Audior"
    ("Office Relocation", r"relocat|address change|registered (address|office)|office transfer|"
                          r"transfer of .*office|name and address"),
    ("Ownership Structure", r"ownership|share ratio|संरचना"),
    ("Board/Management", r"director|chairman|chairperson|board|\bbods?\b|ceo|chief executive|\bmd\b|"
                         r"secr[ea]t[ao]ry|committee|\belection|resignation|appointment|tenure|nominat|oath|"
                         r"sapath|personnel|सञ्चालक|संचालक|कार्यकारी अधिकृत|कम्पनी सचि[वब]|अध्यक्ष|"
                         r"राजीनामा|प्रबन्धक|समिति|adhakchh"),
    ("Investment/Project", r"investment|financial closure|survey licen|land purchase|corporate guarantee|"
                           r"new hydropower|solar project"),
    ("Regulatory/Legal", r"annual fee|नविकरण शुल्क|sebon|\bqiis?\b|बुक बिल्डिङ्ग|बोलपत्र|court|अदालत|"
                         r"कम्पनी ऐन|जवाफ|\bmoa\b|\baoa\b|लाभ वितरण"),
    ("Business Operations", r"logo|branch|migration|सम्झौता|agreement|attorney|\bloan\b|security|quotation"),
    ("Mutual Fund", r"mutual fund|yojana|\bscheme\b|return of equity|maturity"),
    ("Correction", r"correct"),
    # Title carries no topic ("Information [X]", "Company News", "MINUTE", bare ticker "MDB [MDB]").
    ("General Notice", r"^\s*(\d+\w*\s+|regarding |for )?(information|company news|company information|news|"
                       r"notice|letter|minutes?)\b|information letter|attach|जानकारी|"
                       r"^\s*[a-z0-9]+\s*(\[[^\]]*\])?\s*$"),
]

# Optional second level, shown as "Type – Subtype". First match wins; no match keeps the plain type.
MEETING_SUBTYPES = [
    ("Postponed/Changed", r"postpone|exten[st]ion|change in .*date|reschedul|venue|adjourn|delay|cancel|स्थगित"),
    ("Minutes/Decisions", r"minute|miunte|decision|report|appoint|elected|निर्णय|माइन्यूट|माइन्युट|काम कारवाही"),
    ("Completed", r"complet|\bheld\b|successful|conclud|conducted|finished|सम्पन्न"),
    ("Notice", r"notice|announce|intimation|agenda|\bcall|declar|decler|\bdate\b|fixed|organi[sz]|सूचना|बस्ने"),
    ("Meeting Day", r"(agm|sgm|egm|general meeting) of\b"),   # calendar entry: "17th AGM of X Ltd."
]
# Board/Management: "<role> <action>" when a title names both, else the role or action alone.
BOARD_ROLES = [
    ("CEO/Exec", r"\bceo\b|dceo|chief executive|executive director|\bmd\b|managing director|general manager|"
                 r"प्रबन्धक|कार्यकारी अधिकृत"),
    ("Chairman", r"chair|अध्यक्ष|adhakchh"),
    ("Secretary", r"secr[ea]t|सचि[वब]"),
    ("Director", r"director|directos|(?<![a-z])bods?(?![a-z])|board|sanchalak|सञ्चालक|संचालक"),   # "Nepal_BOD"
]
BOARD_ACTIONS = [
    ("Resignation", r"resign|regisn|resigation|\breign|terminat|discontinu|call back|retire|removal|"
                    r"राजीनामा|राजिनामा"),
    ("Tenure", r"tenure|tenture|period|extension|contract|renew|expir|complet"),
    ("Appointment", r"app?[oiu]{1,3}n|nominat|elect|oath|sapath|assum|welcom|\bnew\b|back to role|replace|chang|"
                    r"formation|reconstitut|reform|designation|responsibilit|acting|नियु|नियू|मनोनयन|परिवर्तन|"
                    r"पुनर्गठन|जिम्मेवारी|कायममुकायम|कार्यवाहक"),   # appoint + typos "Appiontment", "Apponited"
]

# IPO: "<investor group> <stage>", e.g. "Foreign Migrants Opening". FPO first: FPO titles also say "general public".
IPO_GROUPS = [
    ("FPO", r"\bfpo\b"),
    ("General Public", r"general public"),
    ("Foreign Migrants", r"foreign|migrant|immigrant"),
    ("Project Locals", r"locals"),
    ("Mutual Fund", r"face value each|yojana|\bfund\b|scheme"),
]
IPO_STAGES = [
    ("Opening", r"opening day"),
    ("Closing", r"closing day"),
    ("Listing", r"\blisting\b"),
]
# Auction: "<share kind> <stage>", e.g. "Promoter Right Bid Opening". Bid Opening before Opening ("Bid Opening Day").
AUCTION_KINDS = [
    ("Promoter Right", r"promoter right"),
    ("Ordinary Right", r"right share"),
    ("Promoter", r"promoter"),
]
AUCTION_STAGES = [
    ("Bid Opening", r"bid opening"),
    ("Opening", r"opening day"),
    ("Closing", r"closing day"),
    ("Cut-off Price", r"cut-?\s*off"),
    ("Listing", r"\blisting\b"),
]


def both(firsts, seconds):
    """Rules labelled "<first> <second>" that match only titles containing both patterns."""
    return [(f"{a} {b}", rf"^(?=.*(?:{a_pat}))(?=.*(?:{b_pat}))") for a, a_pat in firsts for b, b_pat in seconds]


SUBTYPE_RULES = {
    "Board/Management": [
        ("Death/Vacancy", r"death|demise|vacan|रिक्त"),
        ("Committee", r"committee|लेखापरिक्षण समिति"),
        *both(BOARD_ROLES, BOARD_ACTIONS),
        ("Board Meeting", r"meeting|बैठक|decision"),
        *BOARD_ROLES,
        *BOARD_ACTIONS,
    ],
    "IPO": [*both(IPO_GROUPS, IPO_STAGES), *IPO_STAGES, *IPO_GROUPS],
    "Auction": [*both(AUCTION_KINDS, AUCTION_STAGES), *AUCTION_STAGES, *AUCTION_KINDS],
    "AGM": MEETING_SUBTYPES,
    "SGM": MEETING_SUBTYPES,
    "Dividend": [
        ("No Dividend", r"\bno\b.*dividend|not (to )?declare"),
        ("Approval", r"approv|स्वीकृत"),
        ("Proposed", r"propos|recommend|प्रस्तावित"),
        ("Bonus Share", r"bonus"),
        ("Declaration", r"de[cl]{2}[ae]*r|cash dividend"),   # also "Decleration", "Delcaration", "Declration"
    ],
    "Book Closure": [
        ("SGM", r"(?<![a-z])(sgm|egm|eogm)(?![a-z])|special\s+general|extra\s*-?ordinary"),
        ("AGM", r"(?<![a-z])agm(?![a-z])|annual general|साधारण ?सभा"),
        ("Right Share", r"\bright"),
        ("Dividend/Bonus", r"dividend|bonus|लाभांश"),
        ("Mutual Fund", r"\bfund\b|scheme|yojana|return on"),
    ],
}


def classify_event(title):
    t = title.lower()
    for label, pat in TYPE_RULES:
        if re.search(pat, t):
            for sub, sub_pat in SUBTYPE_RULES.get(label, []):
                if re.search(sub_pat, t):
                    return f"{label} – {sub}"
            return label
    return "Other"


def extract_ticker(title):
    m = re.search(r"\[([A-Z0-9&\. ]+)\]", title)
    return m.group(1).strip() if m else ""


def clean_details(detail, title):
    if not detail:
        return ""
    d = detail
    if d.startswith(title):
        d = d[len(title):]
    d = re.sub(r"\s*\n\s*", " | ", d.strip())
    d = re.sub(r"\s{2,}", " ", d)
    d = re.sub(r"(\s*\|\s*)+", " | ", d).strip(" |")
    return d


def build_df(rows):
    df = pd.DataFrame(rows).drop_duplicates(subset=["Date", "Event", "Link"]).reset_index(drop=True)
    if df.empty:
        return df
    df["DateParsed"] = pd.to_datetime(df["Date"], format="%A, %B %d, %Y", errors="coerce")
    df["Day"] = df["DateParsed"].dt.strftime("%A")
    df["Symbol"] = df["Event"].apply(extract_ticker)
    df["Type"] = df["Event"].apply(classify_event)
    df["Details"] = [clean_details(d, e) for d, e in zip(df["Details"], df["Event"])]
    df = df.sort_values(["DateParsed", "Type", "Event"], na_position="last").reset_index(drop=True)
    return df


# ---------- Formatted Excel export ----------
NAVY, GOLD, LIGHT = "1F3864", "BF9000", "F2F4F8"


def export_to_excel(df, path, period_label=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "Events"

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    base = Font(name="Arial", size=10)

    ws["A1"] = "ShareSansar \u2014 Corporate Event Calendar"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
    ws["A2"] = period_label
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="595959")
    ws.merge_cells("A1:F1")
    ws.merge_cells("A2:F2")

    headers = ["Date", "Day", "Symbol", "Type", "Event", "Details"]
    header_row = 4
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[header_row].height = 20

    r = header_row + 1
    prev_date, band = None, False
    for _, row in df.iterrows():
        if pd.notna(row["DateParsed"]) and row["DateParsed"] != prev_date:
            band = not band
            prev_date = row["DateParsed"]
        fill = PatternFill("solid", start_color=LIGHT) if band else None

        date_cell = ws.cell(row=r, column=1,
                            value=row["DateParsed"].to_pydatetime() if pd.notna(row["DateParsed"]) else row["Date"])
        if pd.notna(row["DateParsed"]):
            date_cell.number_format = "dd-mmm-yyyy"
        ws.cell(row=r, column=2, value=row["Day"] if pd.notna(row["Day"]) else "")
        ws.cell(row=r, column=3, value=row["Symbol"])
        ws.cell(row=r, column=4, value=row["Type"])
        ev = ws.cell(row=r, column=5, value=row["Event"])
        if isinstance(row["Link"], str) and row["Link"].startswith("http"):
            ev.hyperlink = row["Link"]
            ev.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        ws.cell(row=r, column=6, value=row["Details"])

        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            if cell.font is None or not cell.font.underline:
                cell.font = base
            if fill:
                cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (5, 6)),
                                       horizontal="center" if c in (1, 2, 3, 4) else "left")
        r += 1

    for col, w in {"A": 13, "B": 11, "C": 10, "D": 42, "E": 55, "F": 70}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{header_row + 1}"
    last_row = r - 1
    ws.auto_filter.ref = f"A{header_row}:F{last_row}"

    # Summary sheet with live COUNTIF formulas
    sm = wb.create_sheet("Summary")
    sm["A1"] = "Event Summary"
    sm["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
    sm["A2"] = period_label
    sm["A2"].font = Font(name="Arial", size=10, italic=True, color="595959")
    sm["A4"] = "Total events"
    sm["A4"].font = Font(name="Arial", size=10, bold=True)
    sm["B4"] = f"=COUNTA(Events!E{header_row + 1}:E{last_row})"
    sm["B4"].font = base

    sm["A6"] = "By Type"
    sm["A6"].font = Font(name="Arial", size=11, bold=True, color=GOLD)
    for c, h in ((1, "Type"), (2, "Count")):
        cell = sm.cell(row=7, column=c, value=h)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=NAVY)
    rr = 8
    for t in sorted(df["Type"].unique()):
        sm.cell(row=rr, column=1, value=t).font = base
        f = sm.cell(row=rr, column=2, value=f'=COUNTIF(Events!D{header_row + 1}:D{last_row},A{rr})')
        f.font = base
        rr += 1

    sm.cell(row=rr + 1, column=1, value="By Date").font = Font(name="Arial", size=11, bold=True, color=GOLD)
    hr2 = rr + 2
    for c, h in ((1, "Date"), (2, "Count")):
        cell = sm.cell(row=hr2, column=c, value=h)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=NAVY)
    rr = hr2 + 1
    for d in df["DateParsed"].dropna().drop_duplicates().sort_values():
        dc = sm.cell(row=rr, column=1, value=d.to_pydatetime())
        dc.number_format = "dd-mmm-yyyy"
        dc.font = base
        f = sm.cell(row=rr, column=2, value=f'=COUNTIF(Events!A{header_row + 1}:A{last_row},A{rr})')
        f.font = base
        rr += 1

    sm.column_dimensions["A"].width = 44
    sm.column_dimensions["B"].width = 10
    wb.save(path)
    return path


def send_mail(subject, body, attachment_path=None):
    if not (GMAIL_USER and GMAIL_APP_PW and MAIL_TO):
        print("Mail skipped: secrets not all set.")
        return
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = GMAIL_USER
    msg["To"] = MAIL_TO
    msg.set_content(body)

    if attachment_path and os.path.exists(attachment_path):
        with open(attachment_path, "rb") as f:
            data = f.read()
        msg.add_attachment(
            data,
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=os.path.basename(attachment_path),
        )

    ctx = ssl.create_default_context()
    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, context=ctx) as s:
        s.login(GMAIL_USER, GMAIL_APP_PW)
        s.send_message(msg)
    print(f"Email sent to {MAIL_TO}")


def main():
    end_date = START_DATE + timedelta(days=NUM_DAYS - 1)
    target = f"{START_DATE:%a %d %b %Y}"
    period_label = f"{START_DATE:%B %d, %Y}"
    file_dates = f"{START_DATE:%Y-%m-%d}"
    if end_date != START_DATE:
        target += f" – {end_date:%a %d %b %Y}"
        period_label += f" – {end_date:%B %d, %Y}"
        file_dates += f" to {end_date:%Y-%m-%d}"
    print(f"Target event dates: {target}")

    rows = scrape()
    df = build_df(rows)
    n = 0 if df.empty else len(df)
    print(f"Scraped {n} events for {target}")

    filename = f"Events {file_dates}.xlsx"
    out_path = os.path.join(SAVE_DIR, filename)

    if df.empty:
        subject = f"[ShareSansar Events] {target} — no events"
        body = (f"No corporate events listed for {target} "
                f"(checked {datetime.now(NPT):%Y-%m-%d %H:%M} NPT).\n"
                f"This is normal for holidays/weekends or if ShareSansar hasn't posted yet.")
        send_mail(subject, body, None)
        return

    export_to_excel(df, out_path, period_label)
    if ALSO_SAVE_CSV:
        df.drop(columns=["DateParsed"]).to_csv(out_path.replace(".xlsx", ".csv"), index=False)

    by_type = df["Type"].value_counts()
    type_lines = "\n".join(f"  {t:<42} {c}" for t, c in by_type.items())
    subject = f"[ShareSansar Events] {target} — {n} event(s)"
    body = (
        f"Corporate events for {target}\n"
        f"(pulled {datetime.now(NPT):%Y-%m-%d %H:%M} NPT, 3 days back)\n\n"
        f"Total: {n}\n"
        f"By type:\n{type_lines}\n\n"
        f"Full formatted workbook attached."
    )
    send_mail(subject, body, out_path)


if __name__ == "__main__":
    main()
